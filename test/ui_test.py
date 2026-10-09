import asyncio, json, pathlib
from playwright.async_api import async_playwright

HERE = pathlib.Path(__file__).parent.resolve()
results = []
def check(name, ok, extra=""):
    results.append(ok); print(("PASS " if ok else "FAIL ") + name + (f"  [{extra}]" if extra else ""))

async def main():
    async with async_playwright() as p:
        b = await p.chromium.launch()
        pg = await b.new_page(viewport={"width": 1280, "height": 900})
        errs = []
        pg.on("pageerror", lambda e: errs.append(str(e)))
        await pg.goto(f"file://{HERE}/harness.html")

        # signed-out screen
        check("login button shown", await pg.locator("#login").count() == 1)
        await pg.click("#login")
        check("login handler fires", await pg.evaluate("window.calls.login") == 1)

        # denied screen
        await pg.evaluate("ui.denied({email:'someone@x.com', provider:'saml.okta'})")
        t = await pg.inner_text("#root")
        check("denied screen shows email+provider", "someone@x.com" in t and "saml.okta" in t)

        # app: active only by default
        counts = await pg.evaluate("window.__counts")
        await pg.evaluate("ui.app({email:'user@example.com'}); ui.setActive(window.__active)")
        n = await pg.locator(".card").count()
        check("active sites listed, closed hidden", n == counts["active"], f"{n} cards")
        check("count label", (await pg.inner_text(".count")).startswith(f"{counts['active']} site"))

        # XSS
        pw = await pg.evaluate("window.pwned")
        check("hostile data not executed", pw is None, f"pwned={pw}")
        check("hostile data shown as literal text", "<img src=x" in await pg.inner_text(".list"))
        check("no <script>/<img> injected from data", await pg.locator(".list script, .list img").count() == 0)

        # search
        await pg.fill("#q", "comcast")
        n2 = await pg.locator(".card").count()
        check("search narrows results", 0 < n2 < n, f"{n2} for 'comcast'")
        await pg.fill("#q", "zzzz-nothing")
        check("empty state", await pg.locator(".empty").count() == 1)
        await pg.fill("#q", "")

        # secondary-only filter
        await pg.check("#sec")
        n3 = await pg.locator(".card").count()
        check("secondary ISP filter", 0 < n3 < n, f"{n3}")
        await pg.uncheck("#sec")

        # expand detail
        first = pg.locator(".card").first
        await first.locator(".head").click()
        check("card expands", await first.locator(".detail").is_visible())
        txt = await first.locator(".detail").inner_text()
        check("detail has Primary ISP + Contacts", "Primary ISP" in txt and "Contacts" in txt.replace("CONTACTS","Contacts") or "CONTACTS" in txt)
        await first.screenshot(path=str(HERE / "shot_card.png"))

        # ---- export (Excel + CSV) of exactly what is on screen ----
        import openpyxl, csv, tempfile
        await pg.fill("#q", "comcast")
        shown = await pg.locator(".card").count()
        async with pg.expect_download() as dl:
            await pg.click("#export-xlsx")
        d = await dl.value
        check("xlsx file name", d.suggested_filename.startswith("branch-isp-info-") and d.suggested_filename.endswith(".xlsx"), d.suggested_filename)
        xp = tempfile.mktemp(suffix=".xlsx"); await d.save_as(xp)
        ws = openpyxl.load_workbook(xp).active
        check("xlsx has header + one row per shown site", ws.max_row == shown + 1, f"{ws.max_row - 1} rows vs {shown} shown")
        check("xlsx header bold + first col is Site", ws["A1"].value == "Site" and ws["A1"].font.b)
        async with pg.expect_download() as dl2:
            await pg.click("#export-csv")
        d2 = await dl2.value
        cp = tempfile.mktemp(suffix=".csv"); await d2.save_as(cp)
        rows_csv = list(csv.reader(open(cp, encoding="utf-8-sig")))
        check("csv has header + one row per shown site", len(rows_csv) == shown + 1, f"{len(rows_csv) - 1} rows")
        await pg.fill("#q", "")

        # ---- edit an existing site ----
        card = pg.locator(".card").first
        await card.locator(".head").click()
        await card.locator(".edit-btn").click()
        check("edit dialog opens", await pg.locator("dialog.modal[open]").count() == 1)
        await pg.fill("dialog [name=description]", "Edited description")
        await pg.click("dialog .crow .link")  # remove first contact
        await pg.click("dialog >> text=+ Add contact")
        await pg.locator("dialog [name=cname]").last.fill("New Person")
        await pg.locator("dialog [name=cphone]").last.fill("512-555-0199")
        await pg.fill("dialog [name=primary_notes]", "changed note")
        await pg.click("dialog button[type=submit]")
        await pg.wait_for_function("window.saves.length == 1")
        sv = await pg.evaluate("window.saves[0]")
        check("edit sends same id", bool(sv["id"]), sv["id"])
        check("edit sends changed fields", sv["data"]["description"] == "Edited description" and sv["data"]["primary"]["notes"] == "changed note")
        check("edit contacts replaced", sv["data"]["contacts"][-1] == {"name": "New Person", "phone": "512-555-0199"})
        check("edit never sends id/_h/updated fields in data", not any(k in sv["data"] for k in ("id", "_h", "updatedAt", "updatedBy")))
        await pg.wait_for_function("document.querySelectorAll('dialog[open]').length == 0")
        check("list shows edited text after save", "Edited description" in await pg.inner_text(".list"))

        # ---- required name + permission error ----
        await pg.click("#add")
        await pg.fill("dialog [name=siteName]", "")
        await pg.evaluate("document.querySelector('dialog form').noValidate = true")
        await pg.click("dialog button[type=submit]")
        check("site name required", "required" in (await pg.inner_text("dialog [role=alert]")).lower())
        await pg.fill("dialog [name=siteName]", "ZZ New Site")
        await pg.evaluate("window.failSave = true")
        await pg.click("dialog button[type=submit]")
        await pg.wait_for_function("document.querySelector('dialog [role=alert]').textContent.includes('permission')")
        check("permission-denied shown, dialog stays open", await pg.locator("dialog[open]").count() == 1)
        await pg.evaluate("window.failSave = false")

        # ---- add a new site ----
        await pg.click("dialog button[type=submit]")
        await pg.wait_for_function("document.querySelectorAll('dialog[open]').length == 0")
        added = await pg.evaluate("window.saves[window.saves.length - 1]")
        check("new site has no id yet (assigned on save)", added["id"] is None)
        await pg.fill("#q", "ZZ New Site")
        check("new site appears in list", await pg.locator(".card").count() == 1)
        await pg.fill("#q", "")

        # ---- editing a site to Closed moves it out of the active list ----
        await pg.fill("#q", "ZZ New Site")
        await pg.locator(".card").first.locator(".head").click()
        await pg.locator(".card").first.locator(".edit-btn").click()
        await pg.select_option("dialog [name=status]", "Closed")
        await pg.click("dialog button[type=submit]")
        await pg.wait_for_function("document.querySelectorAll('dialog[open]').length == 0")
        check("closing a site hides it (closed toggle off)", await pg.locator(".card").count() == 0)
        await pg.fill("#q", "")

        # hostile data stays inert inside the edit form too
        await pg.fill("#q", "img src=x")
        hc = pg.locator(".card").first
        await hc.locator(".head").click(); await hc.locator(".edit-btn").click()
        check("hostile name shown as literal text in form", "<img src=x" in await pg.input_value("dialog [name=siteName]"))
        await pg.click("dialog >> text=Cancel")
        check("no script ran from the form", await pg.evaluate("window.pwned") is None)
        await pg.fill("#q", "")

        # include closed -> lazy load once
        await pg.check("#closed")
        await pg.wait_for_function(f"document.querySelectorAll('.card').length == {counts['active'] + counts['closed']}")
        check("include closed loads all", True, f"{counts['active'] + counts['closed']} cards")
        check("closed loaded once", await pg.evaluate("window.calls.closedLoads") == 1)
        await pg.uncheck("#closed"); await pg.check("#closed")
        check("closed not re-fetched", await pg.evaluate("window.calls.closedLoads") == 1)

        # sign out
        await pg.click("#logout")
        check("logout handler fires", await pg.evaluate("window.calls.logout") == 1)

        # screenshots desktop + mobile
        await pg.uncheck("#closed")
        await pg.screenshot(path=str(HERE / "shot_desktop.png"))
        await pg.set_viewport_size({"width": 390, "height": 800})
        await pg.screenshot(path=str(HERE / "shot_mobile.png"))
        scroll_w = await pg.evaluate("document.documentElement.scrollWidth")
        check("no horizontal scroll on phone width", scroll_w <= 391, f"scrollWidth={scroll_w}")
        check("no page errors", not errs, "; ".join(errs))
        await b.close()
    print(f"\n{sum(results)}/{len(results)} passed")

asyncio.run(main())
